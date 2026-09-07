import docx
import os
import glob
import json
import shutil
import sys
sys.stdout.reconfigure(encoding='utf-8')

def table_to_markdown(table):
    rows = []
    for r in table.rows:
        row_cells = [c.text.replace('\n', ' ').strip() for c in r.cells]
        # Remove duplicate adjacent cells caused by merged cells
        cleaned = []
        for cell in row_cells:
            if not cleaned or cell != cleaned[-1]:
                cleaned.append(cell)
        rows.append(cleaned)
    if not rows:
        return ""
    
    max_cols = max(len(r) for r in rows)
    # Pad rows
    for r in rows:
        while len(r) < max_cols:
            r.append('')
            
    header = "| " + " | ".join(rows[0]) + " |"
    separator = "| " + " | ".join(["---"] * max_cols) + " |"
    data_rows = ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join([header, separator] + data_rows)

def process_docx(file_path):
    doc = docx.Document(file_path)
    file_name = os.path.basename(file_path)
    
    paragraphs = doc.paragraphs
    tables = doc.tables
    
    # Extract metadata
    school = "Trường THPT Khúc Thừa Dụ"
    department = "Tổ Toán - Tin học"
    teacher = "Trần Hoài Thanh"
    topic = ""
    topic_vi = ""
    session_num = 1
    
    # Parse session number from file name
    parts = file_name.replace('.docx', '').split('_')
    if len(parts) > 1 and parts[1].isdigit():
        session_num = int(parts[1])
        
    for p in paragraphs[:12]:
        t = p.text.strip()
        if t.startswith("Teacher’s name:") or t.startswith("Teacher:"):
            teacher = t.split(":", 1)[1].strip()
        elif t.startswith("Topic") or t.startswith("Chủ đề"):
            topic = t
            
    # Vietnamese topic name mapping
    vi_topics = {
        1: "Hệ Bất Phương Trình Bậc Nhất Hai Ẩn & Tam Thức Bậc Hai",
        2: "Hàm Số Lượng Giác & Phương Trình Lượng Giác Chuyên Sâu",
        3: "Dãy Số, Cấp Số Cộng & Cấp Số Nhân",
        4: "Giới Hạn Dãy Số, Giới Hạn Hàm Số & Tính Liên Tục",
        5: "Hàm Số Mũ & Hàm Số Lôgarit, Phương Trình Mũ - Lôgarit",
        6: "Hệ Thức Lượng Trong Tam Giác & Ứng Dụng Hình Học",
        7: "Quan Hệ Song Song & Vuông Góc Trong Không Gian",
        8: "Đại Số Tổ Hợp & Xác Suất Nâng Cao",
        9: "Thống Kê Số Liệu Ghép Nhóm & Các Số Đặc Trưng",
        10: "Các Chuyên Đề Nâng Cao: Nguyên Lý Dirichlet, Bất Biến & Điểm Cực Biên",
        11: "Luyện Đề Thi Thử Số 1 & 2 (Mock Exam 1 – 2)",
        12: "Luyện Đề Thi Thử Số 3 & 4 (Mock Exam 3 – 4)",
        13: "Luyện Đề Thi Thử Số 5 & 6 & Tổng Ôn Toàn Diện (Mock Exam 5 – 6)",
    }
    topic_vi = vi_topics.get(session_num, topic)

    # Document body conversion to clean Markdown
    md_elements = []
    
    # Track tables
    # Since python-docx doesn't provide unified iterator by default, we interleave tables
    # Most lesson plans have 1 table (Vocabulary) near the top after "VOCABULARY"
    table_inserted = False
    
    for p in paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
            
        # Detect vocabulary heading to insert table right after
        if ("VOCABULARY" in txt.upper() or "THUẬT NGỮ" in txt.upper()) and tables and not table_inserted:
            md_elements.append(f"### {txt}\n")
            table_md = table_to_markdown(tables[0])
            if table_md:
                md_elements.append(table_md)
                table_inserted = True
            continue
            
        if p.style.name.startswith('Heading 1'):
            md_elements.append(f"# {txt}\n")
        elif p.style.name.startswith('Heading 2'):
            md_elements.append(f"## {txt}\n")
        elif p.style.name.startswith('Heading 3'):
            md_elements.append(f"### {txt}\n")
        elif txt.startswith(('I.', 'II.', 'III.', 'IV.', 'V.', 'VI.')):
            md_elements.append(f"\n## {txt}\n")
        elif txt.startswith(('1.', '2.', '3.', '4.', '5.', '6.', '7.', '8.', '9.')) and len(txt) < 80:
            md_elements.append(f"\n### {txt}\n")
        elif txt.startswith('•'):
            clean_item = txt.lstrip('•').strip()
            md_elements.append(f"- {clean_item}")
        elif txt.startswith('- '):
            md_elements.append(txt)
        else:
            md_elements.append(txt)
            
    # If table was not inserted and exists
    if not table_inserted and tables:
        for t in tables:
            t_md = table_to_markdown(t)
            if t_md:
                md_elements.append("\n" + t_md + "\n")
                
    markdown_content = "\n\n".join(md_elements)
    
    # Extract vocabulary list
    vocab_list = []
    if tables:
        t0 = tables[0]
        for row in t0.rows[1:]:
            cells = [c.text.replace('\n', ' ').strip() for c in row.cells]
            if len(cells) >= 3 and cells[0]:
                vocab_list.append({
                    "term_en": cells[0],
                    "pronunciation": cells[1] if len(cells) > 1 else "",
                    "term_vi": cells[2] if len(cells) > 2 else ""
                })

    return {
        "id": f"giaoan-{session_num:02d}",
        "session_number": session_num,
        "title": f"Buổi {session_num}: {topic_vi}",
        "title_en": topic,
        "topic": topic_vi,
        "file_name": file_name,
        "file_size": f"{os.path.getsize(file_path) // 1024} KB",
        "teacher": teacher,
        "school": school,
        "department": department,
        "vocabulary": vocab_list,
        "content_markdown": markdown_content
    }

def main():
    docx_files = sorted(glob.glob('giaoan/*.docx'))
    print(f"Processing {len(docx_files)} lesson plan documents...")
    
    # Copy files to public/giaoan/
    os.makedirs('public/giaoan', exist_ok=True)
    for f in docx_files:
        shutil.copy2(f, os.path.join('public/giaoan', os.path.basename(f)))
    print("Copied files to public/giaoan/")
    
    sessions = []
    for f in docx_files:
        data = process_docx(f)
        sessions.append(data)
        print(f"Processed: {data['title']} ({len(data['content_markdown'])} chars, {len(data['vocabulary'])} vocab)")
        
    ts_content = "/**\n * 13 BÀI GIẢNG / GIÁO ÁN BỒI DƯỠNG HSG TOÁN THPT BẰNG TIẾNG ANH\n"
    ts_content += " * Nguồn: Thầy Trần Hoài Thanh - Trường THPT Khúc Thừa Dụ\n"
    ts_content += " */\n\n"
    ts_content += "export interface GiaoAnSession {\n"
    ts_content += "  id: string;\n"
    ts_content += "  session_number: number;\n"
    ts_content += "  title: string;\n"
    ts_content += "  title_en: string;\n"
    ts_content += "  topic: string;\n"
    ts_content += "  file_name: string;\n"
    ts_content += "  file_size: string;\n"
    ts_content += "  teacher: string;\n"
    ts_content += "  school: string;\n"
    ts_content += "  department: string;\n"
    ts_content += "  vocabulary: { term_en: string; pronunciation: string; term_vi: string }[];\n"
    ts_content += "  content_markdown: string;\n"
    ts_content += "}\n\n"
    ts_content += "export const GIAO_AN_SESSIONS: GiaoAnSession[] = " + json.dumps(sessions, ensure_ascii=False, indent=2) + ";\n"
    
    with open('src/data/giaoanData.ts', 'w', encoding='utf-8') as out:
        out.write(ts_content)
        
    print("Successfully generated src/data/giaoanData.ts")

if __name__ == '__main__':
    main()
